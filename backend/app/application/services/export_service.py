"""Exportación a Excel, PDF y CSV."""
from __future__ import annotations

import csv
import io
from datetime import datetime, timezone
from pathlib import Path
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.core.config import settings

LOGO = Path(__file__).resolve().parents[2] / "assets" / "logo-aducorp.png"
NAVY, AZUL, CELESTE = "10203C", "1367AC", "70A9CE"

COLUMNAS_TICKETS = [
    ("numero", "Número"), ("fecha_creacion", "Fecha creación"), ("solicitante", "Solicitante"), ("cargo", "Cargo"), ("area", "Área"),
    ("cliente", "Cliente"), ("correo", "Correo"), ("telefono", "Teléfono"), ("categoria", "Categoría"), ("subcategoria", "Subcategoría"),
    ("prioridad", "Prioridad"), ("responsable", "Responsable"), ("estado", "Estado"), ("asunto", "Asunto"),
    ("fecha_primera_respuesta", "Primera respuesta"), ("fecha_resolucion", "Fecha resolución"), ("sla_respuesta_h", "SLA respuesta (h)"),
    ("sla_resolucion_h", "SLA resolución (h)"), ("sla_resolucion_vence", "Vence"), ("sla_consumido_pct", "SLA consumido %"),
    ("cumplimiento_sla", "Cumplimiento SLA"),
]
COLUMNAS_PDF = ["numero", "fecha_creacion", "solicitante", "cliente", "categoria", "prioridad", "responsable", "estado", "sla_consumido_pct", "cumplimiento_sla"]


def _local(d: datetime | None) -> str:
    if d is None:
        return ""
    return d.replace(tzinfo=timezone.utc).astimezone(ZoneInfo(settings.ZONA_HORARIA)).strftime("%d/%m/%Y %H:%M")


def fila_ticket(t) -> dict:
    return {
        "numero": t.numero, "fecha_creacion": _local(t.fecha_creacion), "solicitante": t.solicitante_nombre, "cargo": t.solicitante_cargo or "",
        "area": t.area.nombre if t.area else "", "cliente": t.cliente.razon_social if t.cliente else "", "correo": t.correo,
        "telefono": t.telefono or "", "categoria": t.categoria.nombre, "subcategoria": t.subcategoria.nombre, "prioridad": t.prioridad.nombre,
        "responsable": t.responsable.nombre if t.responsable else "", "estado": t.estado_rel.nombre, "asunto": t.asunto,
        "fecha_primera_respuesta": _local(t.fecha_primera_respuesta), "fecha_resolucion": _local(t.fecha_resolucion),
        "sla_respuesta_h": round(t.sla_respuesta_min / 60, 2), "sla_resolucion_h": round(t.sla_resolucion_min / 60, 2),
        "sla_resolucion_vence": _local(t.sla_resolucion_vence), "sla_consumido_pct": float(t.sla_consumido_pct or 0),
        "cumplimiento_sla": {"CUMPLIDO": "Cumplido", "INCUMPLIDO": "Incumplido", "EN_CURSO": "En curso", "EXCLUIDO": "Excluido"}[t.cumplimiento_sla],
    }


# ------------------------------------------------------------------ Excel
def _hoja(ws, titulo: str, columnas: list[tuple[str, str]], filas: list[dict], subtitulo: str = "") -> None:
    ws.title = titulo[:31]
    ws["A1"] = f"ADUCORP · Mesa de Servicios — {titulo}"
    ws["A1"].font = Font(bold=True, size=14, color=NAVY)
    ws["A2"] = subtitulo or f"Generado el {datetime.now(ZoneInfo(settings.ZONA_HORARIA)):%d/%m/%Y %H:%M}"
    ws["A2"].font = Font(italic=True, color="5B6B85")
    encabezado = 4
    for i, (_, nombre) in enumerate(columnas, 1):
        c = ws.cell(row=encabezado, column=i, value=nombre)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for r, fila in enumerate(filas, encabezado + 1):
        for i, (clave, _) in enumerate(columnas, 1):
            ws.cell(row=r, column=i, value=seguro(fila.get(clave)))
    for i, (clave, nombre) in enumerate(columnas, 1):
        largo = max([len(str(nombre))] + [len(str(f.get(clave) or "")) for f in filas[:200]])
        ws.column_dimensions[get_column_letter(i)].width = min(max(10, largo + 2), 60)
    ws.freeze_panes = ws.cell(row=encabezado + 1, column=1)
    if filas:
        ws.auto_filter.ref = f"A{encabezado}:{get_column_letter(len(columnas))}{encabezado + len(filas)}"


def seguro(v):
    """Evita inyección de fórmulas al abrir exportaciones en Excel (CSV/XLSX)."""
    if isinstance(v, str) and v[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + v
    return v


def excel(hojas: list[tuple[str, list[tuple[str, str]], list[dict]]], subtitulo: str = "") -> bytes:
    wb = Workbook()
    for i, (titulo, columnas, filas) in enumerate(hojas):
        ws = wb.active if i == 0 else wb.create_sheet()
        _hoja(ws, titulo, columnas, filas, subtitulo)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# -------------------------------------------------------------------- CSV
def csv_bytes(columnas: list[tuple[str, str]], filas: list[dict]) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=";")   # ; para Excel en configuración regional es-EC
    w.writerow([n for _, n in columnas])
    for f in filas:
        w.writerow([seguro(f.get(c, "")) for c, _ in columnas])
    return ("﻿" + buf.getvalue()).encode("utf-8")   # BOM para que Excel respete tildes


# -------------------------------------------------------------------- PDF
EMPRESA = "ADUCORP Asesores Aduaneros y Logísticos Cía. Ltda."


def nombre_empresa(session) -> str:
    """Nombre de la empresa definido en Configuración (empresa.nombre)."""
    from app.infrastructure.repositories.repositorios import CatalogoRepository

    return CatalogoRepository(session).configuracion().get("empresa.nombre") or EMPRESA


def pdf(titulo: str, subtitulo: str, secciones: list[tuple[str, list[tuple[str, str]], list[dict]]], resumen: dict | None = None,
        empresa: str = EMPRESA) -> bytes:
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=1.2 * cm, rightMargin=1.2 * cm, topMargin=1.2 * cm, bottomMargin=1.2 * cm,
                            title=titulo, author="ADUCORP Mesa de Servicios")
    st = getSampleStyleSheet()
    st["Title"].textColor = colors.HexColor(f"#{NAVY}")
    pequeño = st["BodyText"].clone("pequeño", fontSize=7, leading=8.5)
    elementos = []
    if LOGO.exists():
        elementos.append(Image(str(LOGO), width=6.5 * cm, height=1.2 * cm, hAlign="LEFT"))
    elementos += [Paragraph(titulo, st["Title"]), Paragraph(subtitulo, st["Italic"]), Spacer(1, 0.3 * cm)]
    if resumen:
        datos = [[Paragraph(f"<b>{k}</b>", pequeño) for k in resumen], [Paragraph(str(v if v is not None else "—"), pequeño) for v in resumen.values()]]
        t = Table(datos, hAlign="LEFT")
        t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8F0F8")), ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#C5D3E3"))]))
        elementos += [t, Spacer(1, 0.4 * cm)]
    for nombre, columnas, filas in secciones:
        elementos.append(Paragraph(nombre, st["Heading3"]))
        datos = [[Paragraph(f"<b>{n}</b>", pequeño) for _, n in columnas]]
        datos += [[Paragraph(str(f.get(c, "") if f.get(c) is not None else ""), pequeño) for c, _ in columnas] for f in filas[:2000]]
        t = Table(datos, repeatRows=1)
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(f"#{NAVY}")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F4F6FA")]),
            ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#C5D3E3")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ]))
        elementos += [t, Spacer(1, 0.4 * cm)]
        if len(filas) > 2000:
            elementos.append(Paragraph(f"Se muestran 2000 de {len(filas)} filas. Use Excel o CSV para el detalle completo.", pequeño))

    def pie(canvas, documento):
        canvas.saveState()
        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(colors.HexColor("#5B6B85"))
        canvas.drawString(1.2 * cm, 0.7 * cm, f"{empresa} · Documento generado por la Mesa de Servicios")
        canvas.drawRightString(landscape(A4)[0] - 1.2 * cm, 0.7 * cm, f"Página {documento.page}")
        canvas.restoreState()

    doc.build(elementos, onFirstPage=pie, onLaterPages=pie)
    return buf.getvalue()
